"""Generate the Unit 5 Assignment Word document."""
import os
import glob
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
SCREENSHOT_DIR = os.path.join(SCRIPT_DIR, 'Screenshots')

doc = Document()

# ── Styles ──
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)
style.paragraph_format.space_after = Pt(6)

# ── Title page ──
for _ in range(6):
    doc.add_paragraph()
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('IN404 Unit 5 Assignment: Support Vector Machine')
run.bold = True
run.font.size = Pt(16)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run('Shubham Raj').font.size = Pt(14)

subtitle2 = doc.add_paragraph()
subtitle2.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle2.add_run('Purdue University Global').font.size = Pt(14)

subtitle3 = doc.add_paragraph()
subtitle3.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle3.add_run('IN404 – Machine Learning').font.size = Pt(14)

doc.add_page_break()

# ── Item 1: Code & Execution ──
h = doc.add_heading('1. Python Code and Execution', level=1)

doc.add_paragraph(
    'I completed the Python program (IN404_Unit5.py), which is attached separately. '
    'In this program, I read two CSV datasets (IN404_Unit5and6_1.csv and IN404_Unit5and6_2.csv), '
    'explored both dataframes, converted the Date column from an object to a datetime and then '
    'to a float, defined the feature matrix X and target vector y, performed a train/test split, '
    'and trained a Support Vector Machine (SVM) classifier. '
    'I have included screenshots of the execution output below.'
)

# Insert all screenshots in sorted order
screenshots = sorted(glob.glob(os.path.join(SCREENSHOT_DIR, '*.png')))
for img_path in screenshots:
    doc.add_picture(img_path, width=Inches(6.0))
    last_paragraph = doc.paragraphs[-1]
    last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_page_break()

# ── Item 2: Shapes ──
h = doc.add_heading('2. Shapes of Each Dataset', level=1)

shapes = [
    ('df2', '(1000, 2)',
     'I found that Dataframe 2 has 1,000 rows and 2 columns (Number1 and Number2).'),
    ('df1 after datetime conversion to float', '(1000, 3)',
     'After I added the DATE_NUM column, Dataframe 1 has 1,000 rows and 3 columns '
     '(Number1, Date, and DATE_NUM).'),
    ('X', '(1000, 2)',
     'My feature matrix X contains 1,000 rows and 2 columns (Number1 and DATE_NUM).'),
    ('y', '(1000,)',
     'My target vector y is a one-dimensional Series with 1,000 elements (Number1 from df2).'),
    ('X_train', '(750, 2)',
     'My training feature set has 750 rows and 2 columns.'),
    ('X_test', '(250, 2)',
     'My testing feature set has 250 rows and 2 columns.'),
    ('y_train', '(750,)',
     'My training target set has 750 elements.'),
    ('y_test', '(250,)',
     'My testing target set has 250 elements.'),
]

for name, shape, desc in shapes:
    p = doc.add_paragraph()
    run = p.add_run(f'{name}: ')
    run.bold = True
    p.add_run(f'{shape} — {desc}')

doc.add_page_break()

# ── Item 3: Why Date → datetime → float ──
h = doc.add_heading('3. Why the Date Column Was Converted to Datetime and Then to Float', level=1)

doc.add_paragraph(
    'I found that the Date column in df1 was originally stored as a string (object dtype) in the format '
    '"2/4/20". Since machine learning algorithms, including SVM, require all input features to be '
    'numeric, I could not feed a string directly into the model (Pedregosa et al., 2011).'
)
doc.add_paragraph(
    'I performed the conversion in two steps. First, I used pd.to_datetime() to parse the string into a '
    'proper datetime64 type, which gives Python an understanding of the chronological order '
    'and the actual calendar distances between dates (McKinney, 2017). Second, I converted the datetime to a '
    'float by computing the number of days since the earliest date in the column '
    '(df1[\'Date\'].min()). Dividing by np.timedelta64(1, \'D\') turns the timedelta into a '
    'simple floating-point number of days (0.0, 1.0, 2.0, etc.).'
)
doc.add_paragraph(
    'This two-step approach preserves the meaningful time relationships between the dates '
    '(e.g., February 6 is two days after February 4) while producing a numeric feature that '
    'the SVM algorithm can process. If I had simply label-encoded the dates as arbitrary integers, '
    'the relative distances between dates might not have been preserved correctly.'
)

doc.add_page_break()

# ── Item 4: Train/test split ──
h = doc.add_heading('4. Train/Test Split Values', level=1)

doc.add_paragraph(
    'I used the train_test_split() function from scikit-learn with default parameters '
    '(and random_state=1 for reproducibility). By default, scikit-learn uses a 75% train / '
    '25% test split.'
)
doc.add_paragraph(
    'With 1,000 total samples, I obtained:'
)

bullets = [
    'Training set: 750 samples (75%)',
    'Testing set: 250 samples (25%)',
]
for b in bullets:
    doc.add_paragraph(b, style='List Bullet')

doc.add_paragraph(
    'I confirmed this by the shapes: X_train is (750, 2), X_test is (250, 2), '
    'y_train is (750,), and y_test is (250,).'
)

doc.add_page_break()

# ── Item 5: SVM and kernel learning ──
h = doc.add_heading('5. Support Vector Machine (SVM) and Kernel Learning', level=1)

doc.add_paragraph(
    'From my research, a Support Vector Machine (SVM) is a supervised machine learning algorithm used primarily '
    'for classification tasks, though it can also be applied to regression (Cortes & Vapnik, 1995). '
    'The fundamental idea behind SVM is to find the optimal hyperplane that maximally separates data points '
    'belonging to different classes in the feature space. The data points closest to this decision boundary '
    'are called "support vectors," and they are the critical elements that define the position '
    'and orientation of the hyperplane.'
)
doc.add_paragraph(
    'I learned that kernel learning is what makes SVM especially powerful. When data is not linearly separable '
    'in its original feature space, the kernel trick maps the data into a higher-dimensional '
    'space where a linear separator can be found (Hastie et al., 2009). Common kernel functions include the linear '
    'kernel, polynomial kernel, radial basis function (RBF) kernel (which is the default in '
    'scikit-learn\'s SVC), and the sigmoid kernel. The kernel function computes the similarity '
    'between data points in this higher-dimensional space without explicitly performing the '
    'transformation, making the computation efficient.'
)
doc.add_paragraph(
    'I found that SVM is used for tasks such as text classification, image recognition, bioinformatics, '
    'and any scenario where clear class separation is desired. It works well with '
    'high-dimensional data and is effective even when the number of dimensions exceeds the '
    'number of samples (Pedregosa et al., 2011).'
)

doc.add_page_break()

# ── Item 6: ovo vs ovr ──
h = doc.add_heading('6. SVM Decision Function Shapes: OVO and OVR', level=1)

doc.add_paragraph(
    'I learned that SVM was originally designed for binary (two-class) classification. To handle multi-class '
    'problems (like this assignment, where y has three classes: 1, 2, and 3), two common '
    'strategies are used (Pedregosa et al., 2011):'
)

p = doc.add_paragraph()
run = p.add_run('One-vs-One (OVO): ')
run.bold = True
p.add_run(
    'This strategy builds one classifier for every pair of classes. For k classes, it creates '
    'k × (k − 1) / 2 binary classifiers. For my 3-class problem, that means 3 classifiers '
    '(class 1 vs. 2, class 1 vs. 3, and class 2 vs. 3). Each classifier votes on the class, '
    'and the class with the most votes wins. The decision function shape is '
    '(n_samples, n_classes × (n_classes − 1) / 2). In scikit-learn, OVO is always used '
    'internally as the multi-class strategy regardless of the decision_function_shape setting.'
)

p = doc.add_paragraph()
run = p.add_run('One-vs-Rest (OVR): ')
run.bold = True
p.add_run(
    'This strategy builds one classifier per class, where each classifier separates one class '
    'from all others combined. For k classes, it creates k binary classifiers. The decision '
    'function shape is (n_samples, n_classes). Since version 0.19, OVR is the default '
    'decision_function_shape in scikit-learn and is the recommended setting, which is what '
    'I used in my code: svm.SVC(decision_function_shape=\'ovr\').'
)

doc.add_page_break()

# ── Item 7: SVM score ──
h = doc.add_heading('7. SVM Score', level=1)

doc.add_paragraph(
    'The SVM score produced by my program is 0.371 (37.1%).'
)

doc.add_page_break()

# ── Item 8: What the SVM score means ──
h = doc.add_heading('8. Meaning of the SVM Score as Applied to X and y', level=1)

doc.add_paragraph(
    'I used clf.score(X, y), which returns the mean accuracy of the SVM model on the given data. '
    'It compares the model\'s predictions for every sample in X against the true '
    'labels in y, and reports the proportion that were classified correctly (Pedregosa et al., 2011).'
)
doc.add_paragraph(
    'My score of 0.371 means the model correctly classified only 37.1% of the 1,000 samples. '
    'This is a low accuracy score — barely better than random guessing for a 3-class problem '
    '(which would yield roughly 33.3% accuracy by chance). This tells me that the features in X '
    '(Number1 and DATE_NUM from df1) are not strong predictors of the target variable y '
    '(Number1 from df2). The two datasets may not have a meaningful relationship that the SVM '
    'can learn, which makes sense given that the feature values and target values appear to be '
    'largely independent of each other.'
)
doc.add_paragraph(
    'I also noted that clf.score(X, y) evaluates on the full dataset (all 1,000 '
    'samples), not just the test set. Evaluating on the training data as well can sometimes '
    'inflate the score, but even with that advantage my accuracy remains very low, reinforcing '
    'that the model has not found a useful pattern in the data.'
)

doc.add_page_break()

# ── References ──
h = doc.add_heading('References', level=1)
h.alignment = WD_ALIGN_PARAGRAPH.CENTER

references = [
    'Cortes, C., & Vapnik, V. (1995). Support-vector networks. Machine Learning, 20(3), '
    '273–297. https://doi.org/10.1007/BF00994018',

    'Hastie, T., Tibshirani, R., & Friedman, J. (2009). The elements of statistical learning: '
    'Data mining, inference, and prediction (2nd ed.). Springer. '
    'https://doi.org/10.1007/978-0-387-84858-7',

    'McKinney, W. (2017). Python for data analysis: Data wrangling with Pandas, NumPy, and '
    'IPython (2nd ed.). O\'Reilly Media.',

    'Pedregosa, F., Varoquaux, G., Gramfort, A., Michel, V., Thirion, B., Grisel, O., '
    'Blondel, M., Prettenhofer, P., Weiss, R., Dubourg, V., Vanderplas, J., Passos, A., '
    'Cournapeau, D., Brucher, M., Perrot, M., & Duchesnay, E. (2011). Scikit-learn: Machine '
    'learning in Python. Journal of Machine Learning Research, 12, 2825–2830. '
    'https://jmlr.org/papers/v12/pedregosa11a.html',
]

for ref in references:
    p = doc.add_paragraph(ref)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)

# ── Save ──
output_path = 'IN404_ShubhamRaj_Unit5.docx'
doc.save(output_path)
print(f'Document saved to {output_path}')
