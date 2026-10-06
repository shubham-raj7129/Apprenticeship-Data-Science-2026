"""Generate Unit 5 Seminar Notes as a Word document."""

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

os.chdir(os.path.dirname(os.path.abspath(__file__)))

doc = Document()

# ── Styles ──────────────────────────────────────────────
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.15

for level in range(1, 4):
    hs = doc.styles[f'Heading {level}']
    hs.font.name = 'Times New Roman'
    hs.font.color.rgb = None  # auto/black

doc.styles['List Bullet'].font.name = 'Times New Roman'
doc.styles['List Bullet'].font.size = Pt(12)

# ── Title block ─────────────────────────────────────────
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('Unit 5 Seminar Notes — Support Vector Machines')
run.bold = True
run.font.size = Pt(16)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run('IN404 Machine Learning\n').font.size = Pt(11)
meta.add_run('Student: Shubham Raj\n').font.size = Pt(11)
meta.add_run('Instructor: Dr. Sen Baidya').font.size = Pt(11)

doc.add_paragraph()  # spacer

# ── Helper ──────────────────────────────────────────────
def heading(text, level=2):
    doc.add_heading(text, level=level)

def para(text):
    doc.add_paragraph(text)

def bullet(text):
    doc.add_paragraph(text, style='List Bullet')

def bold_para(label, text):
    p = doc.add_paragraph()
    r = p.add_run(label)
    r.bold = True
    p.add_run(text)

# ═══════════════════════════════════════════════════════
heading('Main Points Discussed', level=1)

# --- Housekeeping ---
heading('Housekeeping and Assignment Overview')
bullet('Dr. Baidya opened by checking in with the class and reminding us to review the announcement section for this week\'s coverage.')
bullet('For the seminar alternative assignment, those not attending live need to watch the recording and submit notes per the rubric.')
bullet('This week\'s assignment requires downloading two CSV files (IN404_Unit5and6_1.csv and IN404_Unit5and6_2.csv) and the Python file. All files should be kept in the same folder to avoid path errors.')
bullet('The assignment has 8 items: Item 1 is running the code and taking screenshots, Items 2–8 are conceptual questions based on the output.')
bullet('Dr. Baidya emphasized that results must be printed to a text file and submitted alongside the Word document and .py file — not just screenshots.')

# --- Classification recap ---
heading('Recap of Classification and Supervised Learning')
bullet('Dr. Baidya briefly reviewed what we\'ve covered in previous units: classification is sorting things into predefined groups (spam vs. not spam, cat vs. dog).')
bullet('Nobody writes rules by hand — we collect labeled examples, show them to the model, and the model learns the rule. This is supervised learning.')
bullet('We\'ve been doing binary classification (yes/no, pass/fail) in past units.')
bullet('He used a simple example: two students with features "hours studied" and "problems done" — these are the uppercase X (features), and the result column is the lowercase y (label/target).')

# --- Vectors ---
heading('What is a Vector?')
bullet('Each row of a dataset is one example, and the features of that row form a list of numbers — that list is a vector.')
bullet('For example, row 1 has [4, 10], row 2 has [1, 2], row 3 has [5, 12]. These are all vectors.')
bullet('This is where the "Vector" in Support Vector Machine comes from — nothing fancier than a list of numbers.')

# --- Perceptron ---
heading('The Perceptron')
para('The perceptron is one of the oldest and simplest machine learning models. It draws one straight-line boundary between two groups. Everything on one side gets one label, everything on the other side gets the other.')

bold_para('How it learns: ', 'It starts with a line in an arbitrary place, tests training examples, nudges the line every time it gets one wrong, and stops when no training examples are misclassified.')

para('Dr. Baidya walked through a diagram: Group A has points at 1 and 2, Group B has points at 6 and 7. The perceptron draws a line at 2.5 — everything below is A, everything above is B. It checks all 4 points, they\'re all correct, so it stops.')

bold_para('Two key limitations: ', '')
bullet('Fragile margins — The line at 2.5 leaves almost no room on the A side. A new reading at 2.6 is much closer to A\'s points, but the model labels it B because it\'s above 2.5. If the line had been shifted right, 2.6 would be correctly classified.')
bullet('Straight lines only — If two groups aren\'t arranged linearly (e.g., one group in a cluster surrounded by the other in a ring), the perceptron can never finish learning — no straight line can separate them.')

para('These two limitations are the reason we moved to SVM.')

# --- SVM ---
heading('Support Vector Machines (SVM)')
para('Dr. Baidya broke down the name: Machine = learns a rule from labeled examples; Vector = a list of numbers; Support = a small number of training points that hold the boundary in place.')

para('Unlike the perceptron, SVM places the boundary in the middle of the widest empty space between two groups. The space on both sides of the boundary is called the margin, and SVM\'s entire job is to make it as wide as possible.')

para('He showed the same Group A / Group B data — the perceptron\'s line was at 2.5 (tight margin), but the SVM line has room on both sides.')

bold_para('Why the widest gap matters:', '')
bullet('More room for new data — slightly off-center points still land on the correct side.')
bullet('Only a few points (the support vectors) decide the boundary, so the model stays stable even when distant points shift.')
bullet('The model is measurable — after learning, you ask it to label rows you have answers for and count how many it got right. That fraction is the SVM score, which is one of the assignment questions.')

bold_para('Limitations of SVM: ', 'Slow on very large datasets, gives labels but not probabilities without extra work, and features need to be on comparable scales since the method is distance-based.')

# --- Kernel Learning ---
heading('Kernel Learning')
para('Dr. Baidya showed a diagram where Group A sits in a small central cluster and Group B is scattered around it in a ring. No straight line can separate them — any line would cut through the ring.')

bold_para('The solution: ', 'Since both groups are circles and differ in distance from the center, we can add a third number (Z) for every row: Z = X² + Y² (squared distance from the middle).')

para('He walked through the math: points close to the center get small Z values (Group A), points far from the center get large Z values (Group B). If we sort by Z alone, we can separate the two groups with a straight line in this higher-dimensional space.')

para('This process is kernel learning — a function that gives the model extra information to build the boundary.')

bold_para('Types of kernels:', '')
bullet('Linear — straight boundary, used when groups are already linearly separable.')
bullet('RBF (Radial Basis Function) — curved boundary based on distance. This is the default kernel when you don\'t specify one, and it\'s what the lab code uses.')

para('The kernel decides what shape the boundary is allowed to be.')

# --- OVR vs OVO ---
heading('OVR vs. OVO (Multi-Class Strategies)')
para('When there are more than two groups, one boundary isn\'t enough.')
bullet('OVR (One vs. Rest) — For each group, train a boundary that separates it from everything else combined. 3 groups = 3 boundaries. One result per group.')
bullet('OVO (One vs. One) — Train a boundary for every pair of groups. 3 groups = 3 pairs (A vs. B, A vs. C, B vs. C). Formula: n_groups × (n_groups − 1) / 2.')
para('In the assignment code, we use: svm.SVC(decision_function_shape=\'ovr\').')

# --- Shapes ---
heading('Understanding Shapes in the Output')
bullet('The shape of a table is (rows, columns).')
bullet('Adding a column keeps row count the same but increases column count by 1 — we see this after converting dates.')
bullet('Pulling a single column out gives a shape with only one number and a trailing comma, like (1000,). Dr. Baidya stressed the comma is not a typo — it means there\'s only one dimension.')
bullet('After a train/test split, the two row counts add back up to the original total.')
bullet('He told us to pay attention to how shapes change throughout the code because several assignment questions ask about them.')

# --- Date Conversion ---
heading('Date Conversion (Object → Datetime → Float)')
para('In the CSV file, dates are stored as text. For models, text has no size or distance — you can\'t subtract one piece of text from another.')
bullet('Step 1: pd.to_datetime(df1[\'Date\']) — tells pandas the text is a calendar date. Before this, the type is "object" (pandas name for text). After, it\'s "datetime."')
bullet('Step 2: (df1[\'Date\'] − df1[\'Date\'].min()) / np.timedelta64(1, \'D\') — finds the earliest date, subtracts it from every row (giving a duration), then divides by one day to convert to a plain count of days as a decimal.')
para('Result: earliest date = 0, next day = 1, two days later = 2, etc. Now the model can use dates as numeric features. This is directly tied to the assignment question asking why the date was converted to datetime and then to float.')

# --- Code Walkthrough ---
heading('Code Walkthrough')
bullet('Dr. Baidya walked through the solution code. The helper function writeFunction toggles between console output (PRINT = True) and file output (PRINT = False).')
bullet('The code loads both CSVs, explores DataFrame 1 and 2 (head, dtypes, shape, etc.), converts dates, defines X from DataFrame 1\'s Number1 and DATE_NUM columns, defines y from DataFrame 2\'s Number1, does a train/test split, and builds the SVM.')
bold_para('Key SVM lines he highlighted:', '')
bullet('clf = svm.SVC(decision_function_shape=\'ovr\') — creates the model but doesn\'t train it yet.')
bullet('clf.fit(X_train, y_train) — this is the only line that actually learns; it finds the boundary with the widest margin.')
bullet('clf.score(X, y) — computes the fraction of correctly classified observations.')
bullet('By default the kernel is RBF. To make it linear, you\'d specify kernel=\'linear\'.')

# ═══════════════════════════════════════════════════════
heading('Classroom Discussion', level=1)
bullet('Manish Singh asked whether the kernel in SVM is the same kernel that interacts with the OS in containers. Dr. Baidya clarified they are completely different concepts — the SVM kernel is related to the decision boundary shape (linear, RBF), not the operating system kernel. Manish acknowledged the distinction.')
bullet('Dr. Baidya checked for questions multiple times throughout. The class was mostly quiet, which he took as a sign that the material was clear.')
bullet('He encouraged everyone to email him with questions rather than staying stuck, emphasizing there are no bad questions.')

# ═══════════════════════════════════════════════════════
heading('Points of Interest and Personal Reflections', level=1)
bullet('I found the perceptron-to-SVM progression really helpful for understanding why SVM exists. Seeing the fragile margin problem with concrete numbers (the 2.5 line misclassifying 2.6) made it click why maximizing the margin matters.')
bullet('The kernel learning explanation with the circle-inside-a-ring example was intuitive — adding Z = X² + Y² to create a separable dimension was a clever visual. It made me realize that feature engineering and kernel choice are closely related ideas.')
bullet('The date conversion pipeline (text → datetime → float) is something I can see being useful in many real-world datasets, not just this assignment. It\'s a good general preprocessing pattern to remember.')
bullet('I\'m curious about how SVM performs compared to the logistic regression models we built in Unit 4, especially on text data. The limitations Dr. Baidya mentioned (slow on large datasets, no probabilities) make me wonder when to choose SVM over logistic regression in practice.')
bullet('The distinction between OVR and OVO was new to me. I want to explore whether the choice between them significantly affects accuracy on multi-class problems.')

# ── Save ────────────────────────────────────────────────
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "IN404_ShubhamRaj_Unit5_Seminar.docx")
doc.save(out)
print(f"Saved → {out}")
